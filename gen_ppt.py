"""Generate I4I633_PS-I04.pptx - GreenVision.AI hackathon deck (16:9, dark theme)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- constants
BG = RGBColor(0x1A, 0x1A, 0x2E)
CARD_BG = RGBColor(0x24, 0x24, 0x3E)
GREEN = RGBColor(0x3F, 0xA3, 0x4D)
BLUE = RGBColor(0x4F, 0xA8, 0xD8)
AMBER = RGBColor(0xFF, 0xB7, 0x4D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xE6, 0xE6, 0xE6)
MUTED = RGBColor(0xAA, 0xAA, 0xCC)
BORDER = RGBColor(0x3A, 0x3A, 0x5C)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

TEAM_LINE = "Team ID: I4I 633   |   Team: Hacksmith   |   Leader: Nayana N   |   PS ID: PS-I04 (Software)"
COLLEGE = "J.S.S. Academy of Technical Education, Bengaluru"

OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "submission")
OUT_PATH = os.path.join(OUT_DIR, "I4I633_PS-I04.pptx")

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def add_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    return slide


def add_text(slide, x, y, w, h, text, size=18, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Segoe UI",
             spacing=1.0):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        f = r.font
        f.size = Pt(size)
        f.color.rgb = color
        f.bold = bold
        f.name = font
    return box


def add_bullets(slide, x, y, w, h, items, size=15, color=LIGHT, gap=6):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
        r1 = p.add_run()
        r1.text = "\u25AA  "
        r1.font.size = Pt(size)
        r1.font.color.rgb = GREEN if color is LIGHT else color
        r1.font.bold = True
        r1.font.name = "Segoe UI"
        r2 = p.add_run()
        r2.text = item
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
        r2.font.name = "Segoe UI"
    return box


def add_card(slide, x, y, w, h, fill=CARD_BG, border=BORDER, radius=True):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if border:
        shp.line.color.rgb = border
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_accent_bar(slide, x, y, w=Inches(0.9), h=Pt(5), color=GREEN):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def header(slide, title, accent=GREEN, subtitle=None):
    add_text(slide, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.7),
             title, size=30, color=accent, bold=True)
    add_accent_bar(slide, Inches(0.58), Inches(0.92), Inches(1.4), Pt(5), accent)
    if subtitle:
        add_text(slide, Inches(0.55), Inches(1.02), Inches(12.2), Inches(0.4),
                 subtitle, size=13, color=MUTED)


def footer(slide, idx):
    add_text(slide, Inches(0.45), Inches(7.08), Inches(8), Inches(0.35),
             "GreenVision.AI  |  Hacksmith  |  I4I 633", size=10, color=MUTED)
    add_text(slide, Inches(12.3), Inches(7.08), Inches(0.7), Inches(0.35),
             str(idx), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def chat_card(slide, x, y, w, h, user_msg, bot_msg):
    card = add_card(slide, x, y, w, h)
    add_text(slide, x + Inches(0.18), y + Inches(0.1), w - Inches(0.36),
             Inches(0.32), "USER", size=11, color=BLUE, bold=True)
    add_text(slide, x + Inches(0.18), y + Inches(0.38), w - Inches(0.36),
             Inches(0.62), user_msg, size=11.5, color=WHITE,
             spacing=0.95)
    add_text(slide, x + Inches(0.18), y + Inches(1.02), w - Inches(0.36),
             Inches(0.32), "GREENVISION AI", size=11, color=GREEN, bold=True)
    add_text(slide, x + Inches(0.18), y + Inches(1.3), w - Inches(0.36),
             Inches(0.95), bot_msg, size=11.5, color=LIGHT, spacing=0.95)
    return card


# ================================================================ SLIDE 1
s = add_slide()
add_accent_bar(s, Inches(0), Inches(0), SLIDE_W, Pt(8), GREEN)

add_text(s, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.4),
         "GreenVision.AI", size=66, color=GREEN, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Inches(0.8), Inches(2.95), Inches(11.7), Inches(0.7),
         "AI-Powered Environmental Intelligence for Cities, Industries & People",
         size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(0.8), Inches(3.75), Inches(11.7), Inches(0.5),
         "Turning environmental data into green action",
         size=16, color=MUTED, align=PP_ALIGN.CENTER)

info_card = add_card(s, Inches(1.6), Inches(4.9), Inches(10.1), Inches(1.55))
add_text(s, Inches(1.8), Inches(5.08), Inches(9.7), Inches(0.4),
         TEAM_LINE, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(1.8), Inches(5.52), Inches(9.7), Inches(0.4),
         COLLEGE, size=13, color=LIGHT, align=PP_ALIGN.CENTER)
add_text(s, Inches(1.8), Inches(5.95), Inches(9.7), Inches(0.4),
         "Smart India Hackathon 2026  |  Problem Statement Category: Software",
         size=12, color=MUTED, align=PP_ALIGN.CENTER)

# ================================================================ SLIDE 2
s = add_slide()
header(s, "PROBLEM & SOLUTION")

add_text(s, Inches(0.55), Inches(1.25), Inches(6.0), Inches(0.45),
         "THE PROBLEM", size=19, color=GREEN, bold=True)
add_bullets(s, Inches(0.55), Inches(1.75), Inches(5.9), Inches(3.6), [
    "Indian cities lose 1.5-3% tree cover per decade (ISFR 2023)",
    "No real-time tool to measure urban green cover from imagery",
    "Municipalities plan plantations without data-driven priority mapping",
    "Industrial sites lack green buffer compliance assessment",
    "Citizens cannot assess or act on local environmental data",
], size=14.5, gap=8)

add_text(s, Inches(6.85), Inches(1.25), Inches(6.0), Inches(0.45),
         "OUR SOLUTION", size=19, color=BLUE, bold=True)
add_bullets(s, Inches(6.85), Inches(1.75), Inches(5.95), Inches(3.6), [
    "AI-GIS platform: upload imagery -> environmental diagnosis",
    "3 distinct workspaces: Municipal / Industrial / Citizen",
    "ML models: canopy segmentation, tree detection, scene classification",
    "Location-conditioned species engine (Open-Meteo weather/AQI/soil)",
    "Conversational AI assistant grounded in real data",
], size=14.5, gap=8, color=LIGHT)

metrics = [("7.92%", "Canopy Cover"), ("156", "Trees Detected"),
           ("1,026", "Trees Needed"), ("2.2 t/yr", "CO2 Offset"),
           ("35/35", "Tests Passed")]
mw = Inches(2.42)
mx = Inches(0.55)
my = Inches(5.55)
for i, (val, lab) in enumerate(metrics):
    x = mx + i * (mw + Inches(0.06))
    add_card(s, x, my, mw, Inches(1.15))
    add_text(s, x, my + Inches(0.14), mw, Inches(0.5), val, size=22,
             color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x, my + Inches(0.68), mw, Inches(0.35), lab, size=11.5,
             color=MUTED, align=PP_ALIGN.CENTER)
footer(s, 2)

# ================================================================ SLIDE 3
s = add_slide()
header(s, "THREE WORKSPACES", subtitle="One platform - three tailored experiences for every stakeholder")

cards = [
    ("MUNICIPAL", GREEN, [
        "Upload city/ward imagery",
        "AI canopy detection & analysis",
        "60% green cover target tracking",
        "Planting gap quantification",
        "GIS map with priority zones",
        "Carbon sequestration & oxygen output",
        "Exportable assessment report",
    ]),
    ("INDUSTRIAL", BLUE, [
        "Site green buffer assessment",
        "Regulatory compliance status",
        "Buffer zone planning support",
        "Pollution-tolerant species advisory",
        "Green investment estimates",
        "ESG reporting metrics",
        "Maintenance cost projection",
    ]),
    ("CITIZEN", AMBER, [
        "Location-aware environmental profile",
        "Interactive Leaflet map",
        "Live weather / AQI / soil data",
        "AI chatbot advisor",
        "Natural conversation interface",
        "Green Points rewards system",
        "Community leaderboard",
    ]),
]
cw = Inches(4.05)
cx0 = Inches(0.45)
cy = Inches(1.65)
ch = Inches(5.15)
for i, (title, accent, feats) in enumerate(cards):
    x = cx0 + i * (cw + Inches(0.12))
    add_card(s, x, cy, cw, ch)
    add_accent_bar(s, x, cy, cw, Pt(6), accent)
    add_text(s, x + Inches(0.2), cy + Inches(0.16), cw - Inches(0.4),
             Inches(0.45), title, size=17, color=accent, bold=True,
             align=PP_ALIGN.CENTER)
    add_bullets(s, x + Inches(0.22), cy + Inches(0.72), cw - Inches(0.42),
                ch - Inches(0.9), feats, size=12.5, gap=5)
footer(s, 3)

# ================================================================ SLIDE 4
s = add_slide()
header(s, "GREENVISION AI - CONVERSATIONAL ASSISTANT",
       subtitle="Not a keyword matcher. A context-aware environmental advisor.")

chat_data = [
    ('"Hi, my name is Nayana"',
     '"Hey Nayana! I\'m GreenVision, your environmental assistant. '
     'Ask me about green cover, trees, or your area\'s air quality."'),
    ('"What trees are good for my area?"',
     '"Based on your location in Bengaluru and its climate/soil profile, '
     'top match: 1. Neem (82/100) - drought-tolerant, air-purifying..."'),
    ('"Give me another option"',
     '"Here are other options suited to Bengaluru: 1. Pongamia (75/100) '
     '- hardy canopy native; 2. Indian Laurel (73/100)..."'),
    ('"How many trees do I need?"',
     '"Your ward currently has 156 trees (7.92% canopy). About 1,026 '
     'additional trees to reach the 60% target."'),
]
cw = Inches(6.05)
chh = Inches(2.32)
positions = [(Inches(0.45), Inches(1.6)), (Inches(6.75), Inches(1.6)),
             (Inches(0.45), Inches(4.05)), (Inches(6.75), Inches(4.05))]
for (cx, cyy), (u, b) in zip(positions, chat_data):
    chat_card(s, cx, cyy, cw, chh, u, b)

cap = add_card(s, Inches(0.45), Inches(6.5), Inches(12.43), Inches(0.5),
               fill=CARD_BG, border=None)
add_text(s, Inches(0.6), Inches(6.57), Inches(12.1), Inches(0.4),
         "Session memory   |   Name-aware greetings   |   Follow-up context   |   "
         "Species exclusion   |   Cost follow-ups   |   Honest fallback",
         size=13, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
footer(s, 4)

# ================================================================ SLIDE 5
s = add_slide()
header(s, "ORGANIZATION ECOSYSTEM", subtitle="Every contribution is attributed - individuals build reputations, organizations climb rankings")

add_text(s, Inches(0.55), Inches(1.5), Inches(6.0), Inches(0.45),
         "CONTRIBUTION SYSTEM", size=18, color=BLUE, bold=True)
add_bullets(s, Inches(0.55), Inches(2.0), Inches(5.9), Inches(3.4), [
    "6 org types: Individual, Company, College, School, NGO, Community",
    "Organization stored with each contribution record",
    "Backend aggregates points by organization",
    "4 action types earn Green Points: plant, adopt, verify, review",
], size=14, gap=8)

add_text(s, Inches(6.85), Inches(1.5), Inches(6.0), Inches(0.45),
         "LEADERBOARD", size=18, color=GREEN, bold=True)
add_bullets(s, Inches(6.85), Inches(2.0), Inches(5.95), Inches(3.4), [
    "Individual rankings with podium animation for top 3",
    'Organization aggregation: e.g. Infosys (24 participants, 8,700 pts)',
    "Dedicated tabs: Companies, Colleges, NGOs, Communities",
    "Transparent scoring visible to all participants",
], size=14, gap=8)

note = add_card(s, Inches(0.55), Inches(5.7), Inches(12.23), Inches(0.85))
add_text(s, Inches(0.8), Inches(5.86), Inches(11.8), Inches(0.6),
         '"Green Points are a participation score, not a CO2/O2 measurement."',
         size=15, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
footer(s, 5)

# ================================================================ SLIDE 6
s = add_slide()
header(s, "TECHNICAL ARCHITECTURE", subtitle="End-to-end pipeline from imagery to actionable insight")

flow = ["Frontend\nReact + Leaflet", "Flask API\nREST backend",
        "ML Pipeline\nYOLOv8 + U-Net + Scene Classifier", "Open-Meteo API\nWeather / AQI / Soil"]
fw = Inches(2.7)
fy = Inches(1.7)
fh = Inches(1.15)
for i, label in enumerate(flow):
    x = Inches(0.45) + i * (fw + Inches(0.55))
    add_card(s, x, fy, fw, fh, fill=CARD_BG, border=GREEN if i % 2 == 0 else BLUE)
    parts = label.split("\n")
    add_text(s, x, fy + Inches(0.18), fw, Inches(0.4), parts[0], size=15,
             color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x, fy + Inches(0.58), fw, Inches(0.5), parts[1], size=11.5,
             color=MUTED, align=PP_ALIGN.CENTER)
    if i < len(flow) - 1:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + fw + Inches(0.07),
                                fy + Inches(0.42), Inches(0.42), Inches(0.3))
        ar.fill.solid()
        ar.fill.fore_color.rgb = AMBER
        ar.line.fill.background()
        ar.shadow.inherit = False

stack = [
    ("FRONTEND", GREEN, "React 19  |  Vite  |  TailwindCSS  |  React-Leaflet"),
    ("BACKEND", BLUE, "Flask  |  Python 3.11  |  PyTorch 2.2"),
    ("ML MODELS", GREEN, "YOLOv8 (tree detection)  |  U-Net (canopy segmentation)  |  Scene classifier"),
    ("DATA", BLUE, "Open-Meteo (weather, AQI, soil, elevation)  |  Nominatim (geocoding)"),
    ("STORAGE", GREEN, "Flat JSON (demo-grade)  |  localStorage for identity"),
]
sy = Inches(3.35)
for i, (name, accent, detail) in enumerate(stack):
    y = sy + i * Inches(0.68)
    add_card(s, Inches(0.45), y, Inches(2.2), Inches(0.56))
    add_text(s, Inches(0.45), y + Inches(0.1), Inches(2.2), Inches(0.4), name,
             size=13, color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_card(s, Inches(2.75), y, Inches(10.1), Inches(0.56))
    add_text(s, Inches(2.95), y + Inches(0.1), Inches(9.8), Inches(0.4),
             detail, size=13, color=LIGHT)
footer(s, 6)

# ================================================================ SLIDE 7
s = add_slide()
header(s, "IMPACT & FEASIBILITY")

imp = add_card(s, Inches(0.45), Inches(1.4), Inches(12.43), Inches(2.6))
add_accent_bar(s, Inches(0.45), Inches(1.4), Inches(12.43), Pt(6), GREEN)
add_text(s, Inches(0.75), Inches(1.58), Inches(11.8), Inches(0.45), "IMPACT",
         size=18, color=GREEN, bold=True)
add_bullets(s, Inches(0.75), Inches(2.08), Inches(11.9), Inches(1.8), [
    "Helps municipalities understand environmental health of every neighborhood",
    "Identifies areas with low green cover, high heat, and poor air quality",
    "Enables data-driven plantation decisions - right place, right species",
    "Creates community engagement through Green Champions and leaderboards",
], size=13.5, gap=4)

fea = add_card(s, Inches(0.45), Inches(4.25), Inches(12.43), Inches(2.6))
add_accent_bar(s, Inches(0.45), Inches(4.25), Inches(12.43), Pt(6), BLUE)
add_text(s, Inches(0.75), Inches(4.43), Inches(11.8), Inches(0.45), "FEASIBILITY",
         size=18, color=BLUE, bold=True)
add_bullets(s, Inches(0.75), Inches(4.93), Inches(11.9), Inches(1.8), [
    "Uses readily available satellite/drone imagery and open datasets",
    "Built on scalable, proven AI, GIS, and web technologies",
    "No specialized hardware dependency - runs on commodity infrastructure",
    "Supports phased deployment from ward level to full city scale",
], size=13.5, gap=4)
footer(s, 7)

# ================================================================ SLIDE 8
s = add_slide()
header(s, "FUTURE ROADMAP & REFERENCES")

add_text(s, Inches(0.55), Inches(1.4), Inches(6.0), Inches(0.45),
         "FUTURE ROADMAP", size=18, color=GREEN, bold=True)
add_bullets(s, Inches(0.55), Inches(1.95), Inches(5.9), Inches(4.2), [
    "Pollutant dispersion model for industrial mode",
    "Historical trend analysis (multi-temporal comparison)",
    "Account system with verified identity",
    "Mobile app for field verification",
    "Integration with municipal GIS databases",
], size=14.5, gap=10)

add_text(s, Inches(6.85), Inches(1.4), Inches(6.0), Inches(0.45),
         "REFERENCES", size=18, color=BLUE, bold=True)
add_bullets(s, Inches(6.85), Inches(1.95), Inches(5.95), Inches(4.2), [
    "Smart Cities Mission (MoHUA)",
    "ISRO Bhuvan Geoportal",
    "Forest Survey of India (ISFR)",
    "Open-Meteo API",
    "YOLOv8 / Ultralytics",
], size=14.5, gap=10)

end = add_card(s, Inches(0.55), Inches(6.15), Inches(12.23), Inches(0.7))
add_text(s, Inches(0.7), Inches(6.27), Inches(11.9), Inches(0.5),
         "GreenVision.AI - Turning environmental data into green action.",
         size=16, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
footer(s, 8)

# ---------------------------------------------------------------- save
os.makedirs(OUT_DIR, exist_ok=True)
prs.save(OUT_PATH)
print("Saved:", OUT_PATH)
print("Slides:", len(prs.slides.__iter__.__self__._sldIdLst))
